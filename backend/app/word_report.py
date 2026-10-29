"""Render the current result snapshot as an editable, annotated Word report."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Annotated, Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, Field, model_validator

Probability = Annotated[float, Field(ge=0, le=1, allow_inf_nan=False)]
Text = Annotated[str, Field(max_length=60_000)]
VERDICT = {"low": "较低", "uncertain": "不确定", "high": "较高"}
CONFIDENCE = {"low": "低", "medium": "中", "high": "高"}
SHADE = {"low": "E5EFE8", "uncertain": "FFF1CC", "high": "F8DAD7"}
DISCLAIMER = "本报告是本机模型结果快照，不是学校、期刊、知网或 Turnitin 出具的文件。分数不代表 AI 所写字数占比，也不能证明文字来源。"
REVIEW = {"reviewed": "已复核", "pending": "待确认", "revise": "需修改"}


class ReportModel(BaseModel):
    model_config = ConfigDict(extra="ignore")

    @model_validator(mode="after")
    def xml_text(self):
        for value in self.__dict__.values():
            if isinstance(value, str) and any(
                ord(ch) < 32
                and ch not in "\t\r\n"
                or 0xD800 <= ord(ch) <= 0xDFFF
                or ord(ch) in {0xFFFE, 0xFFFF}
                for ch in value
            ):
                raise ValueError("文字中包含 Word 无法显示的控制字符")
        return self


class Review(ReportModel):
    state: Literal["reviewed", "pending", "revise"]
    note: Annotated[str, Field(max_length=2000)] = ""
    updated_at: Annotated[str, Field(max_length=80)]


class Sentence(ReportModel):
    text: Text
    score: Probability | None = None
    verdict: Literal["low", "uncertain", "high"] | None = None
    review: Review | None = None


class Paragraph(ReportModel):
    index: Annotated[int, Field(ge=0, le=10_000, strict=True)]
    text: Text | None = None
    excerpt: Text = ""
    score: Probability
    verdict: Literal["low", "uncertain", "high"]
    sentences: Annotated[list[Sentence], Field(max_length=1000)] | None = None
    sentenceNote: Annotated[str, Field(max_length=2000)] | None = None


class Ignored(ReportModel):
    text: Text
    reason: Annotated[str, Field(max_length=500)]


class DetectionScope(ReportModel):
    source_name: Annotated[str, Field(max_length=500)]
    source_sha256: Annotated[str, Field(pattern=r"^[a-f0-9]{64}$")]
    total_blocks: Annotated[int, Field(ge=1, le=2_000_000, strict=True)]
    selected_blocks: Annotated[int, Field(ge=1, le=2_000_000, strict=True)]
    total_chars: Annotated[int, Field(ge=0, le=20_000_000, strict=True)]
    selected_chars: Annotated[int, Field(ge=0, le=20_000_000, strict=True)]
    excluded: Annotated[list[Ignored], Field(max_length=2000)] = []
    note: Annotated[str, Field(max_length=2000)]

    @model_validator(mode="after")
    def consistent(self):
        if self.selected_blocks > self.total_blocks or self.selected_chars > self.total_chars:
            raise ValueError("检测范围计数不一致")
        return self


class Calibration(ReportModel):
    status: Literal["default", "candidate", "unvalidated", "validated"]
    applied: bool
    low: Probability
    high: Probability
    note: Annotated[str, Field(max_length=4000)]


class LanguageReport(ReportModel):
    lang: Literal["en", "zh"]
    score: Probability
    verdict: Literal["low", "uncertain", "high"]
    confidence: Literal["low", "medium", "high"]
    char_count: Annotated[int, Field(ge=0, le=2_000_000, strict=True)]
    word_count: Annotated[int, Field(ge=0, le=2_000_000, strict=True)]
    model_id: Annotated[str, Field(max_length=500)]
    model_revision: Annotated[str, Field(max_length=200)] | None = None
    calibration: Calibration | None = None
    explanation: Annotated[str, Field(max_length=8000)] = ""
    basis: Annotated[str, Field(max_length=8000)] = ""
    reading: Annotated[str, Field(max_length=8000)] = ""
    disclaimer: Annotated[str, Field(max_length=4000)] = ""
    paragraphs: Annotated[list[Paragraph], Field(min_length=1, max_length=1000)]
    ignored: Annotated[list[Ignored], Field(max_length=1000)] = []
    detection_scope: DetectionScope | None = None

    @model_validator(mode="after")
    def bounded_text(self):
        if len({p.index for p in self.paragraphs}) != len(self.paragraphs):
            raise ValueError("段落编号不能重复")
        total = sum(
            len(p.text or p.excerpt) + sum(len(s.text) for s in p.sentences or [])
            for p in self.paragraphs
        )
        total += sum(len(i.text) for i in self.ignored)
        total += (
            sum(len(i.text) for i in self.detection_scope.excluded) if self.detection_scope else 0
        )
        total += sum(
            len(s.review.note) for p in self.paragraphs for s in p.sentences or [] if s.review
        )
        if total > 500_000:
            raise ValueError("报告正文过长，请分开导出")
        return self


class SectionError(ReportModel):
    error: Annotated[str, Field(max_length=200)]
    message: Annotated[str, Field(max_length=2000)]


class BilingualReport(ReportModel):
    lang: Literal["bi"]
    note: Annotated[str, Field(max_length=4000)] = ""
    disclaimer: Annotated[str, Field(max_length=4000)] = ""
    en: LanguageReport | SectionError | None = None
    zh: LanguageReport | SectionError | None = None
    detection_scope: DetectionScope | None = None

    @model_validator(mode="after")
    def correct_language(self):
        for lang in ("en", "zh"):
            part = getattr(self, lang)
            if isinstance(part, LanguageReport) and part.lang != lang:
                raise ValueError("合并报告的语言与部分名称不一致")
        return self


class WordRequest(BaseModel):
    result: Annotated[LanguageReport | BilingualReport, Field(discriminator="lang")]


def _shade(run, color: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    run._element.get_or_add_rPr().append(shading)


def _sentence_band(score: float) -> str:
    return "low" if score < 0.4 else "high" if score >= 0.75 else "uncertain"


def _annotated(doc, paragraph: Paragraph):
    body = paragraph.text if paragraph.text is not None else paragraph.excerpt
    text_parts = []
    cursor = 0
    for sentence in paragraph.sentences or []:
        start = body.find(sentence.text, cursor) if sentence.text else -1
        if start < 0:
            text_parts = []
            break
        text_parts.append((body[cursor:start], None, None))
        text_parts.append(
            (
                sentence.text,
                _sentence_band(sentence.score) if sentence.score is not None else None,
                sentence.review,
            )
        )
        cursor = start + len(sentence.text)
    rendered = doc.add_paragraph()
    if text_parts:
        text_parts.append((body[cursor:], None, None))
        for text, band, review in text_parts:
            run = rendered.add_run(text)
            if band:
                _shade(run, SHADE[band])
            if review and text:
                comment = f"人工复核：{REVIEW[review.state]}\n{review.note or '未填写备注'}\n记录时间：{review.updated_at}"
                doc.add_comment(run, text=comment, author="人工复核", initials="WL")
        if any(s.score is None for s in paragraph.sentences or []):
            doc.add_paragraph("部分句子未评分，未评分文字不着色。", style="Caption")
    else:
        rendered.add_run(body)
        doc.add_paragraph(
            "未完成可对应原文的逐句评分；本段仅展示整段分数，正文不着色。", style="Caption"
        )
    if paragraph.sentenceNote:
        doc.add_paragraph(paragraph.sentenceNote, style="Caption")


def _scope_section(doc, scope: DetectionScope | None):
    if scope is None:
        return
    doc.add_heading("实际检测范围", level=2)
    doc.add_paragraph(
        f"文件：{scope.source_name}\n选中 {scope.selected_blocks} / {scope.total_blocks} 个提取段落，{scope.selected_chars} / {scope.total_chars} 个非空白字符。"
    )
    doc.add_paragraph(scope.note, style="Caption")
    if scope.excluded:
        doc.add_heading("用户排除的文字", level=3)
        for item in scope.excluded:
            doc.add_paragraph(item.reason, style="Caption")
            doc.add_paragraph(item.text)


def _language_section(doc, result: LanguageReport):
    doc.add_heading("中文部分" if result.lang == "zh" else "英文部分", level=1)
    summary = doc.add_paragraph()
    lead = summary.add_run(
        f"AI 倾向分数（模型原始分数） {result.score:.1%} · {VERDICT[result.verdict]}\n"
    )
    lead.bold = True
    lead.font.size = Pt(16)
    summary.add_run(
        f"文本充分度：{CONFIDENCE[result.confidence]} · 非空白字符 {result.char_count} · 词数 {result.word_count}"
    )
    doc.add_paragraph("文本充分度反映长度及各段一致性，不是判断正确的概率。")
    legend = doc.add_paragraph("句子颜色：")
    for band in ("low", "uncertain", "high"):
        run = legend.add_run(f" {VERDICT[band]} ")
        _shade(run, SHADE[band])
    doc.add_paragraph(
        "句子颜色按默认 40%／75% 展示门槛；段落和句子未经过整篇样本校准。颜色不表示已确定作者来源。",
        style="Caption",
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, label in zip(table.rows[0].cells, ("段落", "模型分数", "档位")):
        cell.text = label
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "18243B")
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for paragraph in result.paragraphs:
        cells = table.add_row().cells
        for cell, value in zip(
            cells,
            (f"第 {paragraph.index + 1} 段", f"{paragraph.score:.1%}", VERDICT[paragraph.verdict]),
        ):
            cell.text = value
    doc.add_heading("原文与逐句标注", level=2)
    for paragraph in result.paragraphs:
        doc.add_heading(
            f"第 {paragraph.index + 1} 段 · {paragraph.score:.1%} · {VERDICT[paragraph.verdict]}",
            level=3,
        )
        _annotated(doc, paragraph)
    reviews = [
        (p.index, i, s)
        for p in result.paragraphs
        for i, s in enumerate(p.sentences or [])
        if s.review and s.text and s.text in (p.text or p.excerpt)
    ]
    if reviews:
        doc.add_heading("人工复核记录", level=2)
        doc.add_paragraph(
            "人工备注独立于模型分数；原句上的批注可在 Word 的审阅面板中打开。", style="Caption"
        )
        for pi, si, sentence in reviews:
            doc.add_heading(
                f"第 {pi + 1} 段第 {si + 1} 句 · {REVIEW[sentence.review.state]}", level=3
            )
            doc.add_paragraph(sentence.text)
            doc.add_paragraph(sentence.review.note or "未填写备注")
    if result.ignored:
        doc.add_heading("未参与评分的文字", level=2)
        doc.add_paragraph(
            "以下内容保留在报告中，不计入模型评分。原文件版式、图片和脚注未随此结果快照导出。",
            style="Caption",
        )
        for item in result.ignored:
            doc.add_paragraph(item.reason, style="Heading 3")
            doc.add_paragraph(item.text)
    doc.add_heading("判断方法与局限", level=2)
    doc.add_paragraph(f"检测模型：{result.model_id}")
    doc.add_paragraph(f"权重版本：{result.model_revision or '未记录'}", style="Caption")
    if result.calibration:
        info = result.calibration
        doc.add_paragraph(
            f"整篇阈值：{'当前结果使用样本验证门槛' if info.applied else '沿用默认门槛'}（{info.low:.1%}／{info.high:.1%}）。"
        )
        doc.add_paragraph(info.note)
    else:
        doc.add_paragraph("整篇阈值未记录校准证据；分数仍是模型原始信号。")
    for text in (result.basis or result.explanation, result.reading, result.disclaimer):
        if text:
            doc.add_paragraph(text)


def build_word_report(result: LanguageReport | BilingualReport) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.font.color.rgb = RGBColor.from_string("18243B")
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.4
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.color.rgb = RGBColor.from_string("244BD8" if name == "Heading 2" else "18243B")
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(18 if name != "Title" else 6)
        style.paragraph_format.space_after = Pt(9)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for name in ("Caption", "Subtitle"):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.color.rgb = RGBColor.from_string("596579")
        doc.styles[name].font.size = Pt(9)
        doc.styles[name].element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    header = section.header.paragraphs[0]
    mark = header.add_run("WRITING LENS  /  本机写作复核")
    mark.font.size = Pt(8)
    mark.font.color.rgb = RGBColor.from_string("596579")
    header.paragraph_format.space_after = Pt(8)
    doc.core_properties.title = "本机作文模型检测报告"
    doc.core_properties.author = "Essay AI Detector"
    doc.add_heading("作文模型检测报告", level=0)
    doc.add_paragraph(
        datetime.now(timezone.utc).strftime("导出时间：%Y-%m-%d %H:%M UTC"), style="Subtitle"
    )
    doc.add_paragraph(DISCLAIMER)
    _scope_section(doc, result.detection_scope)
    if isinstance(result, BilingualReport):
        doc.add_paragraph(result.note)
        for lang, label in (("en", "英文部分"), ("zh", "中文部分")):
            if lang == "zh":
                doc.add_page_break()
            part = getattr(result, lang)
            if isinstance(part, LanguageReport):
                _language_section(doc, part)
            else:
                doc.add_heading(label, level=1)
                doc.add_paragraph(
                    part.message if isinstance(part, SectionError) else "该语言没有可评分正文。"
                )
    else:
        _language_section(doc, result)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("WRITING LENS · 本机模型结果快照 · 页 ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string("596579")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
