from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from app.main import app


def result(**updates):
    return {
        "lang": "en",
        "score": 0.55,
        "verdict": "uncertain",
        "confidence": "medium",
        "char_count": 500,
        "word_count": 80,
        "model_id": "fixture-model",
        "model_revision": "sha256:" + "a" * 64,
        "basis": "Fixture basis",
        "reading": "Fixture reading",
        "disclaimer": "Fixture disclaimer",
        "paragraphs": [
            {
                "index": 0,
                "text": "First text.\n\tSecond text.  Unscored end.",
                "score": 0.55,
                "verdict": "uncertain",
                "sentences": [
                    {"text": "First text.", "score": 0.8, "verdict": "high"},
                    {"text": "Second text.", "score": None, "verdict": None},
                ],
            }
        ],
        "ignored": [{"text": "References\nPreserved source", "reason": "参考文献"}],
        **updates,
    }


def test_word_download_preserves_exact_text_and_only_colors_scored_sentences():
    response = TestClient(app).post("/api/reports/word", json={"result": result()})
    assert response.status_code == 200
    assert "wordprocessingml.document" in response.headers["content-type"]
    doc = Document(BytesIO(response.content))
    body = next(p for p in doc.paragraphs if p.text.startswith("First text."))
    assert body.text == result()["paragraphs"][0]["text"]
    assert "F8DAD7" in body.runs[1]._r.xml
    assert "w:shd" not in next(r for r in body.runs if r.text == "Second text.")._r.xml
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Preserved source" in text and "fixture-model" in text
    assert "部分句子未评分" in text
    with ZipFile(BytesIO(response.content)) as archive:
        assert "word/vbaProject.bin" not in archive.namelist()
        assert 'TargetMode="External"' not in archive.read("word/_rels/document.xml.rels").decode()


def test_unmatched_sentence_snapshot_falls_back_to_complete_paragraph():
    data = result()
    data["paragraphs"][0]["sentences"] = [{"text": "Old pre-edit sentence.", "score": 0.9}]
    response = TestClient(app).post("/api/reports/word", json={"result": data})
    doc = Document(BytesIO(response.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert data["paragraphs"][0]["text"] in text
    assert "Old pre-edit sentence." not in text
    assert "未完成可对应原文" in text


def test_bilingual_partial_failure_is_kept_in_merged_word():
    data = {
        "lang": "bi",
        "note": "中英分别计算",
        "en": result(),
        "zh": {"error": "too_short", "message": "中文正文过短"},
    }
    response = TestClient(app).post("/api/reports/word", json={"result": data})
    assert response.status_code == 200
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert "英文部分" in text and "中文部分" in text and "中文正文过短" in text


def test_report_rejects_bad_text_and_mismatched_language_without_server_error():
    client = TestClient(app)
    bad = result()
    bad["paragraphs"][0]["text"] = "Invalid\u0001XML"
    for data in (bad, result(score=2), {"lang": "bi", "zh": result()}, {"lang": "en"}):
        response = client.post("/api/reports/word", json={"result": data})
        assert response.status_code == 400
        assert response.json()["error"] == "invalid_report"


def test_manual_review_is_a_native_word_comment_and_in_the_review_log():
    data = result()
    data["paragraphs"][0]["sentences"][0]["review"] = {
        "state": "reviewed",
        "note": "已对照原始笔记，保留 <引用>。",
        "updated_at": "2026-10-03T12:00:00Z",
    }
    response = TestClient(app).post("/api/reports/word", json={"result": data})
    assert response.status_code == 200
    with ZipFile(BytesIO(response.content)) as archive:
        comments = archive.read("word/comments.xml").decode()
        assert "已复核" in comments and "已对照原始笔记" in comments and "&lt;引用&gt;" in comments
        assert "commentRangeStart" in archive.read("word/document.xml").decode()
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert "人工复核记录" in text and "已对照原始笔记" in text
    assert data["score"] == 0.55


def test_scope_and_excluded_originals_are_in_word_and_bad_reviews_are_rejected():
    data = result(
        detection_scope={
            "source_name": "paper.txt",
            "source_sha256": "a" * 64,
            "total_blocks": 3,
            "selected_blocks": 1,
            "total_chars": 900,
            "selected_chars": 500,
            "excluded": [
                {"text": "Excluded appendix preserved verbatim.", "reason": "用户排除：附录"}
            ],
            "note": "只检测所选正文",
        }
    )
    response = TestClient(app).post("/api/reports/word", json={"result": data})
    assert response.status_code == 200
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert "选中 1 / 3" in text and "Excluded appendix preserved verbatim." in text
    data["paragraphs"][0]["sentences"][0]["review"] = {
        "state": "reviewed",
        "note": "bad\u0001note",
        "updated_at": "now",
    }
    assert TestClient(app).post("/api/reports/word", json={"result": data}).status_code == 400


def test_score_is_labelled_as_a_raw_model_score_not_an_ai_rate():
    """“AI 率”会被读成 AI 所写文字的占比；标签自己要写明是模型原始分数。"""
    response = TestClient(app).post("/api/reports/word", json={"result": result()})
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert "AI 倾向分数（模型原始分数） 55.0%" in text
    assert "AI 率" not in text
