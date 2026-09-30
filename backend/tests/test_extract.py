from io import BytesIO

import pytest
from docx import Document

from app.errors import DetectError
from app.extract import extract_from_bytes


def test_utf8_txt():
    got = extract_from_bytes("a.TXT", "你好世界\n第二段".encode("utf-8"))
    assert got.text == "你好世界\n第二段"
    assert got.natural_paragraphs is None


def test_gbk_txt():
    raw = "中文作业".encode("gbk")
    got = extract_from_bytes("hw.txt", raw)
    assert "中文作业" in got.text


@pytest.mark.parametrize(
    "encoding",
    ["utf-8-sig", "utf-16", "utf-16-be", "utf-32"],
)
@pytest.mark.parametrize(
    "text",
    ["这是第一段。我们去了图书馆。\n\n第二段在这里。", "First paragraph here.\n\nSecond one."],
)
def test_txt_with_bom_decodes_without_a_stray_mark(encoding, text):
    raw = text.encode(encoding)
    if encoding == "utf-16-be":
        raw = b"\xfe\xff" + raw  # 该编码本身不写 BOM，手动补上
    got = extract_from_bytes("notes.txt", raw)
    assert got.text == text
    assert "\ufeff" not in got.text


def test_docx_joins_nonempty_paragraphs():
    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("")
    doc.add_paragraph("第二段")
    buf = BytesIO()
    doc.save(buf)
    got = extract_from_bytes("paper.docx", buf.getvalue())
    assert got.natural_paragraphs == ("第一段", "第二段")
    assert got.text == "第一段\n\n第二段"


def test_rejects_corrupt_pdf():
    with pytest.raises(DetectError) as ei:
        extract_from_bytes("x.pdf", b"%PDF")
    assert ei.value.code == "parse_failed"


def _text_pdf(message: str) -> bytes:
    safe = message.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 24 Tf 72 720 Td ({safe}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


def test_text_pdf_is_extracted():
    got = extract_from_bytes("paper.pdf", _text_pdf("Hello from a text PDF."))
    assert "Hello from a text PDF" in got.text


def test_blank_pdf_is_treated_as_scan():
    from io import BytesIO

    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = BytesIO()
    writer.write(buf)
    with pytest.raises(DetectError) as ei:
        extract_from_bytes("scan.pdf", buf.getvalue())
    assert ei.value.code == "scanned_pdf"


def test_empty_txt_is_empty_error():
    with pytest.raises(DetectError) as ei:
        extract_from_bytes("a.txt", b"   \n")
    assert ei.value.code == "empty"


def test_corrupt_docx():
    with pytest.raises(DetectError) as ei:
        extract_from_bytes("a.docx", b"not-a-zip")
    assert ei.value.code == "parse_failed"


# ---- PDF 重新拼段 ----

from app.extract import pdf_paragraphs  # noqa: E402

_FULL = "This line is written to fill the whole column width of the page body"


def _full(n: int) -> str:
    # 真实正文行不会逐字重复；完全相同的行会被当成页眉页脚。
    return f"{_FULL} {'abcdefgh'[n] * 3}"


def test_pdf_paragraph_spanning_pages_is_joined():
    page1 = f"Course Paper Header\n{_full(1)}\n{_full(2)}\n{_full(3)}\n1"
    page2 = (
        f"Course Paper Header\n{_full(4)}\nand it ends here.\n"
        f"{_full(5)}\n{_full(6)}\nShort last line.\n2"
    )
    paras = pdf_paragraphs([page1, page2])
    # 页眉、页码被去掉；跨页的第一段接成一段；短行收尾。
    assert len(paras) == 2
    assert paras[0].count(_FULL) == 4 and paras[0].endswith("and it ends here.")
    assert paras[1].endswith("Short last line.")
    assert all("Header" not in p for p in paras)


def test_pdf_blank_line_splits_and_page_number_variants_drop():
    page = f"- 3 -\n{_FULL}\n\n{_FULL}\n第 3 页 共 9 页"
    assert pdf_paragraphs([page]) == [_FULL, _FULL]


def test_pdf_hyphen_and_cjk_joins():
    wide = "这一行中文写满了整行版心宽度用来测试拼接是否不插入多余的空格啊"
    assert pdf_paragraphs([f"{wide}\n{wide}\n结束。"]) == [wide + wide + "结束。"]
    en = "The model was trained on a wide range of essays and it was evalu-"
    assert pdf_paragraphs([f"{en}\nated carefully on a held out set of student writing ok\nDone."])[0].startswith(
        en[:-1] + "ated"
    )


def test_pdf_heading_is_its_own_paragraph():
    paras = pdf_paragraphs([f"1. Introduction\n{_FULL}\n{_FULL}\nEnd."])
    assert paras[0] == "1. Introduction"
    assert len(paras) == 2


def test_pdf_word_like_civil_at_top_is_kept():
    paras = pdf_paragraphs([f"Civil\n{_FULL}\nEnd."])
    assert paras[0] == "Civil"
